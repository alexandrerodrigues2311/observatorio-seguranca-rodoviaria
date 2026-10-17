"""Relatório DOCX e versão de leitura na tela a partir das mesmas evidências."""
from pathlib import Path
import json,html,hashlib,datetime
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT,WD_CELL_VERTICAL_ALIGNMENT
R=Path(__file__).resolve().parents[1];D=R/'dist';E=D/'exports'
def read(n):return json.loads((D/'data'/f'{n}.json').read_text(encoding='utf-8'))
def N(x,d=0):return f'{x:,.{d}f}'.replace(',','X').replace('.',',').replace('X','.')
ml=read('aprendizado');ex=read('exposicao');multi=read('multinomial');model=read('modelo');persist=read('persistencia');audit=read('auditoria');sources=read('fontes');gov=json.loads((E/'operacionalizacao-governanca.json').read_text(encoding='utf-8'));feas=json.loads((E/'viabilidade-protocolo.json').read_text(encoding='utf-8'));improvements=json.loads((E/'dez-melhorias-pesquisa.json').read_text(encoding='utf-8'));blocks=[]
def heading(t):blocks.append({'type':'heading','text':t})
def para(t):blocks.append({'type':'paragraph','text':t})
def table(headers,rows):blocks.append({'type':'table','headers':headers,'rows':rows})
def image(name,caption):blocks.append({'type':'image','file':name,'caption':caption})
heading('Objeto e conclusão do relatório')
para('Este relatório documenta o material de apoio à avaliação do projeto O que o Estado vê escolhas de governança e visibilidade nas rodovias federais, de Débora Cristina Coutinho Vilas Bôas, orientado pela Profa. Dra. Suzana Rodrigues no Doutorado em Administração da Universidade FUMEC. A qualificação foi informada para 27 de outubro de 2026 às 15h. As extensões foram preparadas como subsídio à discussão e não são atribuídas como resultados já produzidos pela autora.')
para('O conjunto permite auditar fontes, explorar registros, comparar associações com exposição, examinar gravidade e persistência e selecionar candidatos à análise documental. As análises de aprendizado de máquina acrescentam ganhos preditivos modestos e perfis descritivos estáveis. O mecanismo institucional permanece uma contribuição a demonstrar: dados cadastrais, modelos e resíduos não identificam por si só qualidade de governança, eventos ocultos ou ausência de ação.')
para('A documentação é apresentada na tela e neste arquivo DOCX. Os códigos, fontes, dados derivados e saídas analíticas são públicos. Não foi publicada a íntegra do documento privado da proposta. O relatório é uma síntese metodológica e crítica de apoio, sujeita à discussão com autora e orientação.')
heading('A tese e a base teórica do projeto')
para('A seção 1.5 do projeto argumenta que os registros oficiais dependem da capacidade de observar e registrar do arranjo que os produz e que reconhecer o risco não assegura resposta estatal. Nesta síntese, a tese a demonstrar é que os arranjos de governança condicionam a produção da informação sobre os sinistros e sua transformação em resposta. Essa proposição não equivale a concluir que concessão seja mais ou menos segura.')
para('O projeto já articula legibilidade estatal em Scott, governança regulatória em arranjos híbridos, instrumentos de ação pública, mensuração e auditabilidade, uso da informação de desempenho e valor público. Portanto, essas bases não foram inventadas pelo laboratório. A revisão organiza essa articulação em diagramas: como regras e instrumentos produzem informação, como categorias permitem comparar e quais responsabilidades e recursos conectam o diagnóstico à resposta.')
para('A VGPR é definida no projeto como capacidade de observar e registrar riscos e resultados vinculados ao local, à data e ao responsável. É um construto formativo. Instrumentos de monitoramento e vinculação são dimensões anteriores ao sinistro e medidas principais; cobertura é dimensão resultante e critério de validação. A proposta não adota índice composto principal. A cobertura efetiva ainda precisa de evidência que não seja apenas o resíduo de um modelo.')
para('Transparência trata de divulgação e accountability de explicação e responsabilização; a proposta situa a visibilidade na observação que antecede essas relações. O PDCA organiza a trilha de evidências, mas incentivos, coordenação e capacidade explicam possíveis interrupções. Persistência dos danos não prova interrupção. A contribuição pretendida exige demonstrar que o mecanismo acrescenta explicação distinta a conceitos próximos e admite refutação.')
heading('O que o projeto já previa e o que foi acrescentado')
table(['Componente','Projeto original','Trabalho executado ou proposto no laboratório'],[['Problema e teoria','VGPR, legibilidade, arranjos híbridos e resposta pelo PDCA','Síntese explícita, diagramas e confronto dos limites de identificação'],['Dados','Integração de PRF, SNV, tráfego, instrumentos, intervenções e documentos','Aquisição de 26 arquivos originais, manifesto e derivados rastreáveis'],['Análise quantitativa','Diagnóstico, contagens esperadas, gravidade e persistência','Auditoria, logística, NB2 com exposição, multinomial e testes temporais'],['Aprendizado de máquina','Não é o núcleo analítico original descrito','Boosting de gravidade, boosting Poisson e agrupamento avaliados'],['Governança','Operacionalização prevista por instrumentos, vínculo e documentos','Matriz de viabilidade e contagens cadastrais executadas; histórico institucional pendente'],['Produtos','Base integrada, mapa e modelo de governança observável','Painel interativo, notebook, apresentação web e PowerPoint, relatório DOCX']])
heading('Universos e fontes efetivamente utilizados')
para('PRF 2022–2025 contém 278.103 ocorrências, 23.274 pessoas mortas e 77.521 feridos graves na versão preservada. Uma linha representa uma ocorrência. O total de pessoas mortas difere do número de sinistros fatais. A versão anual pode divergir de números preliminares do projeto por revisões da fonte. A auditoria não corrige divergências por suposição.')
table(['Ano','Sinistros','Mortes','Feridos graves','BR km ou UF inválidos'],[[str(a['ano']),N(a['ocorrencias']),N(a['mortos']),N(a['feridos_graves']),N(a['localizacao_invalida'])] for a in audit])
para('O SNV 202604A é referência cadastral e cartográfica de 2026. A geometria foi simplificada apenas para apresentação. Vinculação por UF, BR e intervalo de km não certifica a exatidão do local nem reconstrói a vigência histórica dos contratos. O mapa do painel utiliza uma projeção Mercator explícita com escala uniforme de tela e a mesma transformação para limites e rodovias. Essa projeção preserva ângulos locais, não áreas; não há estiramento independente para preencher o cartão.')
para('O PNCT disponibiliza estimativas anuais de tráfego para os dois arranjos. Utilizam-se os limites dos segmentos estimados, não a repetição dos limites originais após subdivisão. Somam-se os sentidos C e D; a extensão é a diferença entre limite final e inicial; a exposição multiplica VMD, extensão e dias do ano. Em 2025, o VMD usa julho de 2024 a junho de 2025 segundo a nota técnica. A incerteza do VMD não foi propagada aos intervalos dos modelos.')
table(['Ano','Segmentos elegíveis','PRF vinculados','PRF total','Cobertura'],[[str(a['ano']),N(a['segmentos_elegiveis']),N(a['ocorrencias_vinculadas']),N(a['ocorrencias_prf']),N(100*a['ocorrencias_vinculadas']/a['ocorrencias_prf'],1)+'%'] for a in ex['coverage']])
para('Os arquivos da ANTT incluem instrumentos, investimentos e passagens em pedágios. Passagens não são veículos únicos nem VMD por segmento. Linhas estruturalmente inválidas são segregadas; uma foi identificada no acompanhamento anual. Cadastros de SAT, radares e CFTV têm campos de situação e referência de malha; ano_do_pnv_snv não é ano de instalação. Ausência de data de inativação preenchida não comprova operação contínua.')
heading('Métodos estatísticos executados e interpretação')
para('A auditoria examina IDs por ano, conflitos, datas, localização, coordenadas, classificação e soma de pessoas. Gravidade deriva dos campos numéricos: morte prevalece sobre lesão grave, que prevalece sobre leve. Divergência aritmética sinaliza investigação, sem imputação ou correção automática. Um mesmo registro pode atender a várias regras de qualidade; suas contagens não devem ser somadas como casos exclusivos.')
para('A regressão logística usa o desfecho de ao menos uma morte entre sinistros registrados, com noite, uso urbano, pista dupla ou múltipla, chuva, fim de semana, UF e ano. A amostra é de 277.388 ocorrências após 715 exclusões. A covariância sanduíche é agrupada em 249 corredores UF–BR, com correção finita e quantil t. A OR é uma razão de chances condicionada aos registros, não risco por viagem ou efeito causal.')
table(['Fator','OR','IC 95 por cento'],[[t['term'],N(t['odds_ratio'],3),N(t['low'],3)+' a '+N(t['high'],3)] for t in model['terms'] if t['term'] in ['noite','urbano','dupla','multipla','chuva','fim_semana']])
para('A classificação multinomial distingue quatro categorias exclusivas de gravidade máxima. Referência: sem vítimas. Controles incluem as características do modelo logístico e seis tipos mais frequentes selecionados no treino, além de UF e ano no ajuste completo. A RRR compara a razão entre a probabilidade de uma categoria e a de sem vítimas entre níveis do fator; não é risco relativo marginal. A hipótese IIA e a cobertura distinta da referência são limitações. Há sensibilidade sem tipo de sinistro. O tipo é conhecido após o evento; o modelo não antecipa acidentes antes da viagem.')
para('O multinomial foi treinado em 2022–2023 e avaliado em 2024–2025 sem indicadores futuros de ano. Log loss de '+N(multi['validation']['log_loss'],3)+' versus '+N(multi['validation']['baseline_log_loss'],3)+' da referência; acurácia de '+N(100*multi['validation']['accuracy'],1)+'% versus '+N(100*multi['validation']['baseline_accuracy'],1)+'%. Isso mostra pequena melhora probabilística e quase nenhum ganho na classe mais provável. O arquivo completo contém matriz de confusão e calibração por desfecho.')
para('Os modelos NB2 de sinistros e mortes usam offset da exposição em unidades de 100 milhões de veículos km, com controles de pista, log VMD, proporção de pesados, classificação não simulável, UF e ano. A variância permite sobredispersão. A IRR é associação condicionada ao modelo. A classificação de administração é cadastral, sem validação de cada transição contratual.')
table(['Análise de mortes','IRR de concessão','IC 95 por cento'],[[label,N(t['irr'],3),N(t['low'],3)+' a '+N(t['high'],3)] for label,mod in [('Principal',ex['models']['mortos']),('Sem 2025',ex['sensitivity']['sem_2025']),('Sem não simuláveis',ex['sensitivity']['sem_nao_simulavel'])] for t in mod['terms'] if t['term']=='concessao'])
para('Na sensibilidade sem segmentos não simuláveis, o IC inclui 1. Os dados e o modelo tornam-se compatíveis com ausência de associação nessa escala; isso não comprova igualdade. A comparação ajustada pode ter sentido diferente da bruta. Não se conclui superioridade causal de um arranjo. O teste temporal de mortes teve desvio Poisson '+N(ex['validation']['mean_poisson_deviance'],3)+' e MAE '+N(ex['validation']['mae'],3)+'.')
para('A persistência correlaciona mortes em 2022–2023 e 2024–2025 em 6.389 intervalos UF–BR de 10 km com registro em algum ano. Spearman '+N(persist['rho'],3)+'; IC percentílico '+N(persist['rho_low'],3)+' a '+N(persist['rho_high'],3)+' em 500 reamostragens de corredores, semente 20261017. O universo utiliza presença em todo o período, inclusive anos posteriores ao primeiro biênio; não é um censo da malha nem previsão prospectiva. A persistência ajuda a selecionar locais para investigar resposta, sem provar inação.')
heading('Aprendizado de máquina e protocolo de avaliação')
para('Foram executados três experimentos. Para as tarefas preditivas, dois candidatos previamente definidos são escolhidos por treino em 2022 e validação em 2023; o ajuste final usa 2022–2023 e o teste usa 2024–2025. Não se escolhem parâmetros com os resultados de teste. O teste espacial temporal reserva corredores inteiros por hash com semente 20261017 e também os exclui da seleção. Cerca de 20% dos corredores são reservados, não necessariamente 20% dos registros. Categorias e transformações são aprendidas no treino. O scikit learn utilizado é 1.7.2.')
heading('Classificação por boosting')
c=ml['classifier']['calibrated'];cr=ml['classifier']['reference'];v=multi['validation']
para('HistGradientBoostingClassifier estima quatro categorias com as seis variáveis de contexto, UF e tipos selecionados em 2022. Não usa mortos ou feridos como preditores. O número de árvores, folhas e regularização é escolhido na validação. Uma temperatura escalar é ajustada nas probabilidades de 2023 e aplicada ao ajuste final; a efetividade dessa transferência é avaliada no teste, e as métricas sem calibração também são preservadas. Não há balanceamento artificial de classes.')
table(['Modelo no teste temporal','Log loss menor é melhor','Brier menor é melhor','Acurácia'],[['Constante',N(cr['log_loss'],4),N(cr['brier'],4),N(100*cr['accuracy'],2)+'%'],['Multinomial',N(v['log_loss'],4),N(v['brier_multiclasse'],4),N(100*v['accuracy'],2)+'%'],['Boosting calibrado',N(c['log_loss'],4),N(c['brier'],4),N(100*c['accuracy'],2)+'%']])
para('A melhora em log loss e Brier é modesta e não representa comprovação inferencial de superioridade. A sensibilidade fatal pela categoria de maior probabilidade é '+N(100*c['recall'][3],2)+'%; AP fatal '+N(c['fatal_ap'],3)+' versus prevalência '+N(c['fatal_prevalence'],3)+'. Acurácia balanceada '+N(c['balanced_accuracy'],3)+' e F1 macro '+N(c['macro_f1'],3)+' mostram o limite que a acurácia global esconde. O ganho útil está na comparação de probabilidades e na investigação de padrões; não é ferramenta operacional validada.')
image('ml-classificacao.png','Figura 1 Avaliação temporal da gravidade e sensibilidade por categoria')
heading('Regressão por boosting com exposição')
r=ml['regression']['model'];rr=ml['regression']['reference'];nr=ml['regression']['nb2_reference'];gain=100*(1-r['poisson_deviance']/nr['mean_poisson_deviance'])
para('HistGradientBoostingRegressor usa perda Poisson, alvo de mortes por exposição e pesos proporcionais à exposição. A previsão de taxa é multiplicada pela exposição para voltar à contagem. A parte ajustável da perda corresponde à modelagem de contagens com exposição. Os controles são comparáveis aos do NB2. O VMD contemporâneo está disponível na base avaliada; esse teste não representa previsão que antecede a disponibilidade do tráfego.')
table(['Modelo no mesmo teste','Desvio Poisson','MAE','Mortes previstas'],[['Taxa constante',N(rr['poisson_deviance'],4),N(rr['mae'],4),N(rr['predicted'],1)],['NB2',N(nr['mean_poisson_deviance'],4),N(nr['mae'],4),N(nr['predicted'],1)],['Boosting',N(r['poisson_deviance'],4),N(r['mae'],4),N(r['predicted'],1)]])
para('O teste contém '+N(r['n'])+' segmentos ano e '+N(r['observed'])+' mortes registradas. O boosting reduz o desvio em '+N(gain,2)+'% em relação ao NB2 e prevê '+N(r['predicted'],1)+' mortes. Essa diferença é uma medida preditiva na amostra de teste, sem IC pareado inferencial. As previsões e os resíduos são disponibilizados para auditoria. Um resíduo pode refletir variação aleatória, covariável ausente, erro cadastral, exposição ou especificação; não é uma estimativa identificada de sinistro oculto.')
heading('Teste espacial temporal')
cs=ml['classifier_spatial'];rs=ml['regression_spatial'];table(['Tarefa','Teste reservado','Modelo','Referência constante'],[['Gravidade',N(cs['calibrated']['n'])+' ocorrências','Log loss '+N(cs['calibrated']['log_loss'],4),'Log loss '+N(cs['reference']['log_loss'],4)],['Contagens',N(rs['model']['n'])+' segmentos ano','Desvio '+N(rs['model']['poisson_deviance'],4),'Desvio '+N(rs['reference']['poisson_deviance'],4)]])
para('O teste espacial temporal avalia outros corredores e anos, excluídos do treino e da seleção. Não é validação externa com uma fonte independente de sinistros e não identifica causalidade. As métricas entre testes temporal e espacial não devem ser comparadas como se as amostras tivessem a mesma composição.')
heading('Agrupamento e candidatos documentais')
g=ml['clusters'];para('O agrupamento usa '+N(g['n'])+' segmentos com os mesmos limites físicos em 2022 e 2023. As variáveis são log VMD, proporção de pesados, duplicação e log1p das taxas registradas de sinistros e mortes. RobustScaler e corte em mais ou menos cinco unidades padronizadas reduzem a influência de extremos. Foram avaliados k de 2 a 6 por silhouette e ARI em oito subamostras de 80%. O critério exige grupo mínimo de 5% e ARI médio de pelo menos 0,70 antes de escolher a maior silhouette.')
para('A solução selecionou '+str(g['selected']['k'])+' grupos, silhouette '+N(g['selected']['silhouette'],3)+' e ARI médio '+N(g['selected']['ari_mean'],3)+'. Estabilidade nas reamostragens não comprova que existam classes naturais ou que elas representem governança. Taxas iniciais fazem parte da formação: diferenças iniciais de danos entre grupos não são validação independente.')
table(['Grupo','Segmentos','VMD mediano','Taxa de sinistros','Taxa de mortes'],[['G'+str(p['group']),N(p['segments']),N(p['vmda_median'],1),N(p['crash_rate'],3),N(p['death_rate'],3)] for p in g['profiles']])
para('As taxas são agregadas por exposição em unidades de 100 milhões de veículos km, referentes ao período inicial. O grupo de poucos registros pode refletir perfis de risco, observação ou problemas de compatibilidade; não recebe rótulo de segurança ou de invisibilidade. Na comparação futura, '+N(g['future_matching_segments'])+' de '+N(g['n'])+' segmentos possuem limites exatamente iguais em ao menos um ano posterior, cobertura de '+N(100*g['future_matching_segments']/g['n'],1)+'%. Essa seleção limita a comparação temporal.')
image('ml-agrupamentos.png','Figura 2 Projeção PCA dos perfis iniciais para visualização')
para('Foram listados candidatos documentais com mesma UF e grupo e administração cadastral distinta, próximos nas características padronizadas. A distância não é geográfica nem medida de equivalência institucional. A lista prioriza revisão de episódios e documentos; não constitui pareamento causal. Deve incluir casos contrários à hipótese e confirmar localização, vigência, obrigações e resposta.')
heading('Importância e utilidade para a tese')
para('A importância por permutação mede a perda preditiva ao embaralhar uma variável no teste; não altera o ajuste. Variáveis correlacionadas podem substituir umas às outras. Três repetições indicam variação computacional, sem constituir intervalos inferenciais. A utilidade está em formular perguntas sobre combinações e em verificar se modelos mais flexíveis acrescentam previsão. Ganhos modestos e baixa identificação fatal restringem o uso operacional.')
para('A contribuição institucional não decorre do algoritmo. A seleção de perfis e os casos mal previstos ajudam a escolher episódios para reconstruir quem tinha informação, obrigação e recursos, o que foi decidido, o que foi entregue e o que mudou depois. Só essa ligação entre medida institucional e trajetória temporal pode aproximar a evidência do mecanismo proposto na tese.')
heading('Operacionalização observável da governança')
for item in gov:
 para(item['dimensao']+' — '+item['indicador']+'. '+item['definicao_operacional']+'. Unidade: '+item['unidade']+'. Fonte: '+item['fonte']+'. Situação: '+item['situacao']+'. Condição para teste: '+item['condicao_para_teste']+'.')
heading('Viabilidade de cada etapa do protocolo')
for item in feas:para(item['item']+'. Situação: '+item['situacao']+'. O que já é possível: '+item['o_que_os_dados_permitem']+' O que falta: '+item['o_que_falta'])
heading('Moderação mediação e efeito de intervenções')
para('Moderação institucional é uma hipótese ainda não estimada: a associação do arranjo ou de uma obrigação com resultados pode variar conforme capacidade operacional anterior. A interação exige medir governança e capacidade por trecho e período, verificar suporte comum e apresentar taxas ou probabilidades previstas com incerteza. Interações entre pista e horário disponíveis podem explorar heterogeneidade, mas não substituem essa mensuração.')
para('Mediação institucional também permanece proposta. É necessário observar governança antes da resposta e resultado depois, com medida confiável do mediador e pressupostos sobre confundimento das relações. Produto de coeficientes ou ajuste de equações estruturais não prova mecanismo causal. A série curta, a agregação e a ausência do histórico são limitações materiais.')
para('Intervenções datadas podem permitir estudo de evento ou diferenças em diferenças com comparação adequada, tendências anteriores plausíveis, análise de antecipação e tratamento da adoção em datas distintas. A mera existência de um contrato ou obra não assegura essas condições. Essas análises não foram executadas nesta edição.')
heading('Dez melhorias propostas para o desenvolvimento do projeto')
for i,item in enumerate(improvements,1):para(str(i)+' '+item['melhoria']+'. '+item['contribuicao'])
heading('Cronograma e ponto de alinhamento')
para('O projeto prevê outubro de 2026 a setembro de 2027. A apresentação sintetiza as fases e conserva o Quadro 8 integral em apêndice. A agenda informada pelo avaliador é 27 de outubro de 2026 às 15h; o quadro original marca qualificação em novembro e dezembro. A divergência foi explicitada, sem deslocamento silencioso das atividades.')
raw=json.loads((R/'dist/apresentacao/cronograma-original.json').read_text(encoding='utf-8'));months=['out 2026','nov 2026','dez 2026','jan 2027','fev 2027','mar 2027','abr 2027','mai 2027','jun 2027','jul 2027','ago 2027','set 2027'];table(['Atividade original','Meses previstos'],[[r[0],'; '.join(months[i] for i,x in enumerate(r[1:]) if x)] for r in raw[2:]])
heading('Produtos e rastreabilidade')
links=[('Painel','https://observatorio-seguranca-rodoviaria.vercel.app/'),('Colab','https://colab.research.google.com/github/alexandrerodrigues2311/observatorio-seguranca-rodoviaria/blob/main/notebooks/seguranca_rodoviaria.ipynb'),('GitHub','https://github.com/alexandrerodrigues2311/observatorio-seguranca-rodoviaria'),('Fontes preservadas','https://github.com/alexandrerodrigues2311/observatorio-seguranca-rodoviaria/releases/tag/dados-2026-10-17'),('Apresentação','https://observatorio-seguranca-rodoviaria.vercel.app/apresentacao/')]
for label,url in links:para(label+' '+url)
para('O notebook contém 16 células de código executadas, incluindo o script integral de cada etapa, as verificações, os resultados e as interpretações. As transformações e definições de cada gráfico são documentadas. Os filtros alteram a exploração e os resumos indicados; não reestimam os modelos nacionais. A apresentação principal tem 14 slides e dois apêndices opcionais, com notas, fontes e roteiro de 20 minutos. O PowerPoint preserva texto, diagramas, tabelas e gráfico editáveis; os mapas são ativos cartográficos derivados das fontes.')
para('O pipeline verifica integridade SHA256, totais, cobertura, convergência e coerência dos intervalos e modelos. O multinomial inclui conferência numérica de gradiente e posto da Hessiana. A interface e a apresentação passam por revisão de diagramação, arquivos, filtros, mapas, notas, exportação e acesso público. Esses controles não validam hipóteses causais nem garantem publicação acadêmica.')
heading('Fontes e referências para interpretação')
para('A base teórica e o cronograma foram conferidos na proposta, especialmente seções 1.5, 1.6, 2.1 a 2.8, 3, 4 e Quadro 8. Scott, Espeland e Stevens, Power, Moore e Bozeman são apresentados conforme a articulação da autora. A agenda complementar de identificação deve ser discutida na revisão teórica. As referências abaixo apoiam implementação e aprofundamento, sem transformar associações em causalidade.')
for url in ['https://scikit-learn.org/stable/auto_examples/linear_model/plot_poisson_regression_non_normal_loss.html','https://scikit-learn.org/stable/modules/ensemble.html','https://scikit-learn.org/stable/modules/clustering.html','https://scikit-learn.org/stable/modules/permutation_importance.html','https://imai.fas.harvard.edu/research/mediation/','https://bcallaway11.github.io/did/','https://www.urbandemographics.org/files/2021_Alves_highway-concessions-road-safety.pdf']:para(url)
heading('Manifesto das fontes preservadas')
para('Cada entrada informa arquivo, tamanho, momento de coleta ou verificação, origem e SHA256. O provedor pode alterar o arquivo na URL original. A release preserva a edição utilizada. Data de coleta ou verificação não é necessariamente data de publicação do órgão. Arquivos em cache conservam seu conteúdo; o hash é o identificador da versão. As fontes originais mantêm seus termos de uso.')
for f in sources:
 blocks.append({'type':'source','file':f['file'],'bytes':f.get('bytes',0),'url':f['url'],'sha':f.get('sha256','não disponível'),'time':f.get('retrieved_at','não disponível')})
# Shared structured report is rendered on screen as well as DOCX.
(D/'data/relatorio-conteudo.json').write_text(json.dumps(blocks,ensure_ascii=False,indent=2),encoding='utf-8');(D/'data/relatorio-conteudo.js').write_text('window.RELATORIO_CONTEUDO='+json.dumps(blocks,ensure_ascii=False)+';',encoding='utf-8')
# DOCX in Letter portrait, no decorated title or callout boxes.
doc=Document();sec=doc.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=Inches(.72);sec.bottom_margin=Inches(.72);sec.left_margin=Inches(.75);sec.right_margin=Inches(.75)
for style in ['Normal','Body Text']:
 st=doc.styles[style];st.font.name='Arial';st.font.size=Pt(11);st.paragraph_format.space_after=Pt(7);st.paragraph_format.line_spacing=1.12
for name,size in [('Title',25),('Heading 1',16),('Heading 2',13)]:
 st=doc.styles[name];st.font.name='Arial';st.font.size=Pt(size);st.font.color.rgb=RGBColor(0,0,0);st.paragraph_format.space_before=Pt(15);st.paragraph_format.space_after=Pt(8)
doc.add_paragraph('Relatório das análises e extensões do projeto de governança rodoviária','Title');doc.add_paragraph('Material de apoio à avaliação e ao desenvolvimento da pesquisa');doc.add_paragraph('Edição final de 17 de outubro de 2026');doc.add_paragraph('Projeto de Débora Cristina Coutinho Vilas Bôas\nOrientação Profa. Dra. Suzana Rodrigues\nUniversidade FUMEC\nDoutorado em Administração')
doc.add_paragraph('Este documento detalha o projeto de origem, as análises efetivamente executadas, os resultados e suas limitações, os indicadores de governança e os testes que ainda dependem de integração histórica. A conclusão central é que a extensão melhora a auditabilidade e a seleção de perguntas e casos, mas ainda não identifica o mecanismo institucional proposto.');doc.add_page_break()
header=sec.header.paragraphs[0];header.text='Caderno de apoio à discussão do projeto';header.style=doc.styles['Normal'];header.runs[0].font.size=Pt(9);header.runs[0].font.color.rgb=RGBColor.from_string('666666')
footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT;footer.add_run('Página ');field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
def link(par,url,label=None):
 rel=par.part.relate_to(url,'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',is_external=True);h=OxmlElement('w:hyperlink');h.set(qn('r:id'),rel);r=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'973C32');pr.append(color);r.append(pr);t=OxmlElement('w:t');t.text=label or url;r.append(t);h.append(r);par._p.append(h)
def addtable(b):
 t=doc.add_table(rows=1,cols=len(b['headers']));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
 widths=[7/len(b['headers'])]*len(b['headers']);
 if len(widths)==2:widths=[4.5,2.5]
 if len(widths)==3:widths=[2.1,2.45,2.45]
 for j,h in enumerate(b['headers']):t.rows[0].cells[j].text=h
 for r in b['rows']:
  cells=t.add_row().cells
  for j,v in enumerate(r):cells[j].text=str(v)
 for ri,row in enumerate(t.rows):
  for j,c in enumerate(row.cells):
   c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;tcPr=c._tc.get_or_add_tcPr();sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'333333' if ri==0 else 'F5F5F5' if ri%2==0 else 'FFFFFF');tcPr.append(sh);borders=OxmlElement('w:tcBorders')
   for edge in ['top','left','bottom','right','insideH','insideV']:
    el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
   tcPr.append(borders);mar=OxmlElement('w:tcMar')
   for edge in ['top','left','bottom','right']:
    el=OxmlElement('w:'+edge);el.set(qn('w:w'),'95');el.set(qn('w:type'),'dxa');mar.append(el)
   tcPr.append(mar)
   for par in c.paragraphs:
    par.paragraph_format.space_after=Pt(3);par.paragraph_format.space_before=Pt(3);par.paragraph_format.line_spacing=1.05;par.paragraph_format.keep_with_next=ri==0
    for run in par.runs:run.font.size=Pt(10);run.font.bold=ri==0;run.font.color.rgb=RGBColor(255,255,255) if ri==0 else RGBColor(0,0,0)
  if ri==0:rep=OxmlElement('w:tblHeader');row._tr.get_or_add_trPr().append(rep)
  split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(split)
 doc.add_paragraph().paragraph_format.space_after=Pt(4)
for b in blocks:
 typ=b['type']
 if typ=='heading':doc.add_heading(b['text'],level=1)
 elif typ=='paragraph':
  if b['text'].startswith('http'):link(doc.add_paragraph(),b['text'])
  elif ' https://' in b['text'] and len(b['text'])<300:
   label,url=b['text'].split(' https://',1);p=doc.add_paragraph(label+' ');link(p,'https://'+url)
  else:doc.add_paragraph(b['text'])
 elif typ=='table':addtable(b)
 elif typ=='image':doc.add_picture(str(E/b['file']),width=Inches(6.7));p=doc.add_paragraph(b['caption']);p.paragraph_format.space_after=Pt(12);p.runs[0].font.size=Pt(10)
 elif typ=='source':
  p=doc.add_paragraph(b['file']);p.paragraph_format.keep_with_next=True;p.runs[0].bold=True;p.runs[0].font.size=Pt(10)
  p=doc.add_paragraph(str(b['bytes'])+' bytes · '+b['time']);p.paragraph_format.keep_with_next=True
  for run in p.runs:run.font.size=Pt(10)
  p=doc.add_paragraph();link(p,b['url']);p.paragraph_format.keep_with_next=True
  p=doc.add_paragraph('SHA256 '+b['sha']);p.runs[0].font.name='Courier New';p.runs[0].font.size=Pt(9)
for root in [doc.styles.element,doc.element]:
 for border in list(root.iter(qn('w:pBdr'))):border.getparent().remove(border)
path=E/'relatorio-extensoes-pesquisa.docx';doc.save(path)
# On-screen report, not a JSON-only download.
htmlparts=[]
for b in blocks:
 if b['type']=='heading':htmlparts.append('<h2>'+html.escape(b['text'])+'</h2>')
 elif b['type']=='paragraph':htmlparts.append('<p>'+html.escape(b['text'])+'</p>')
 elif b['type']=='table':htmlparts.append('<div class="table-wrap"><table><thead><tr>'+''.join('<th>'+html.escape(h)+'</th>' for h in b['headers'])+'</tr></thead><tbody>'+''.join('<tr>'+''.join('<td>'+html.escape(str(x))+'</td>' for x in row)+'</tr>' for row in b['rows'])+'</tbody></table></div>')
 elif b['type']=='image':htmlparts.append('<figure><img style="max-width:100%" src="'+b['file']+'"><figcaption>'+html.escape(b['caption'])+'</figcaption></figure>')
 elif b['type']=='source':htmlparts.append('<details><summary>'+html.escape(b['file'])+'</summary><p>'+str(b['bytes'])+' bytes · '+html.escape(b['time'])+'</p><p><a href="'+html.escape(b['url'])+'">'+html.escape(b['url'])+'</a></p><code>SHA256 '+b['sha']+'</code></details>')
(E/'relatorio-extensoes-pesquisa.html').write_text('<!doctype html><html lang="pt-BR"><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1"><title>Relatório das análises e extensões</title><style>body{max-width:1100px;margin:40px auto;padding:20px;font:17px/1.7 system-ui;color:#302b29}h2{margin-top:36px}table{border-collapse:collapse;width:100%}td,th{padding:12px;border:1px solid #ddd;text-align:left;vertical-align:top}.table-wrap{overflow:auto}code{overflow-wrap:anywhere}a{color:#973c32}figure{margin:24px 0}</style><h1>Relatório das análises e extensões do projeto</h1><p><a href="relatorio-extensoes-pesquisa.docx" download>Baixar relatório em DOCX</a></p>'+''.join(htmlparts)+'</html>',encoding='utf-8')
print('DOCX gerado',path.stat().st_size,'bytes;',len(blocks),'blocos documentados')
